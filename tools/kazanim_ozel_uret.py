# -*- coding: utf-8 -*-
"""kazanim_html_uret.py'den SONRA çalışır; TYMM sitesinin yetersiz kaldığı üç durumu resmî kaynaktan düzeltir.

1) Din Kültürü ve Ahlak Bilgisi 5-6-7: TYMM sitesindeki DKAB ünite sayfaları ders kitabı (2026-27), DÖGM yıllık planı ve
   ÖDSGM senaryo tablosuyla ÇELİŞİYOR (ör. DKAB.5.1.2 "müşahede edebilme" ↔ sitede "gözleme dayalı tahmin edebilme").
   Kitap + plan + senaryo birbirini tuttuğu için çıktı ve süreç bileşenleri DÖGM 2026-2027 çerçeve yıllık planından alınır.
2) İngilizce 7: okulda 2026-27'de ESKİ program uygulanır (senaryo ve plan E7.x.y kodlu). Dosya MEB 2026-27 yıllık planından kurulur.
3) Türkçe 5 T.O.5.14: tam program PDF'inde (02.09.2026) 6. Tema'da var, TYMM tema sayfasının çıktı listesinde yok → eklenir.

  python tools/kazanim_ozel_uret.py [--out C:/ttx/out] [--plan-kok .../MEB-KAYNAK/2026-2027/05-yillik-plan]
"""
import argparse, glob, re, datetime
from pathlib import Path
import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("--out", default=r"C:\ttx\out")
ap.add_argument("--plan-kok", default=r"C:\Users\cetin\Projects\MEB-KAYNAK\2026-2027\05-yillik-plan")
a = ap.parse_args()
OUT, PK = Path(a.out), Path(a.plan_kok)
bugun = datetime.date.today().isoformat()
def tr_upper(s): return s.replace("i", "İ").replace("ı", "I").upper()
def temiz(s): return re.sub(r"\s+", " ", str(s or "")).strip()

# ---------------------------------------------------------------- 1) DKAB
for g in (5, 6, 7):
    xl = glob.glob(str(PK / "din-kulturu-ve-ahlak-bilgisi" / str(g) / "*.xlsx"))[0]
    ws = openpyxl.load_workbook(xl, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    # başlık iki satırlı; asıl satırda "SÜREÇ BİLEŞENLERİ" ayrı hücrededir (üst satırda "ÇIKTILARI VE SÜREÇ…" birleşik)
    hi = next(i for i, r in enumerate(rows) if any("SÜREÇ" in temiz(c).upper() and "ÇIKTI" not in temiz(c).upper() for c in r))
    H = [temiz(c).upper() for c in rows[hi]]
    col = lambda key: next(i for i, h in enumerate(H) if key in h)
    c_saat, c_unite, c_konu, c_ak, c_cikti, c_bil = col("DERS"), col("ÜNİTE"), col("KONU"), col("ANAHTAR"), col("ÖĞRENME ÇIKTI"), col("SÜREÇ")
    uniteler = {}; sira = []; cur = None
    for r in rows[hi + 1:]:
        u = temiz(r[c_unite]) if c_unite < len(r) else ""
        if u and not re.search(r"OKUL TEMELL|SOSYAL ETK", u, re.I):
            cur = u
            if cur not in uniteler: uniteler[cur] = dict(saat=0, konu=[], ak=[], cikti={}); sira.append(cur)
        if not cur: continue
        U = uniteler[cur]
        try: U["saat"] += int(float(r[c_saat]))
        except Exception: pass
        for k, lst in ((c_konu, U["konu"]), (c_ak, U["ak"])):
            v = temiz(r[k]) if k < len(r) else ""
            if v and v not in lst: lst.append(v)
        ct = str(r[c_cikti] or "") if c_cikti < len(r) else ""
        bl = str(r[c_bil] or "") if c_bil < len(r) else ""
        m = re.match(rf"\s*(DKAB\.{g}\.\d+\.\d+)\.?\s*(.+)", ct, re.S)
        if m:
            kod = m[1]; ifade = temiz(m[2])
            bil = [temiz(x) for x in re.split(r"\n(?=[a-zçğ]\)\s)", bl.strip()) if temiz(x)]
            if kod not in U["cikti"]: U["cikti"][kod] = [ifade, bil]
            elif len(bil) > len(U["cikti"][kod][1]): U["cikti"][kod][1] = bil
    toplam = sum(len(U["cikti"]) for U in uniteler.values())
    L = [f"# {g}. Sınıf Din Kültürü ve Ahlak Bilgisi — TYMM Öğrenme Çıktısı Verisi (2026-2027)", "",
         f"**Kaynak (doğrulanmış):** Din Öğretimi Genel Müdürlüğü 2026-2027 çerçeve yıllık planı (`{Path(xl).name}`), {bugun} tarihinde aktarıldı "
         "(`tools/kazanim_ozel_uret.py`). Öğrenme çıktıları ve süreç bileşenleri **birebir**; asla yeniden yazılmaz.", "",
         "> **KAYNAK NOTU:** TYMM sitesindeki DKAB ünite sayfalarında bazı çıktı ifadeleri farklıdır (ör. \"gözleme dayalı tahmin edebilme\"). "
         "2026-2027 ders kitabı, DÖGM yıllık planı ve ÖDSGM ortak sınav senaryo tablosu bu dosyadaki ifadeleri kullanır; geçerli olan budur.",
         "> Ortak yazılı sınavlarda geçerli kapsam, ÖDSGM konu-soru dağılım tablosudur (`tymm-acik-uclu-sinav/references/senaryolar/`).", "",
         f"Toplam ünite: {len(sira)} · Öğrenme çıktısı: {toplam} · Ders saati (planda ünitelere düşen): {sum(U['saat'] for U in uniteler.values())}", "",
         "## Ünite tablosu", "", "| Kod/No | Ad | Çıktı | Saat |", "|---|---|---|---|"]
    bloklar = []
    for i, u in enumerate(sira, 1):
        U = uniteler[u]; ad = tr_upper(re.sub(r"^\d+\.\s*", "", u))
        L.append(f"| ÜNİTE {i} | {ad} | {len(U['cikti'])} | {U['saat']} |")
        bloklar += [f"## ÜNİTE {i} — {ad}", "", f"**Ders saati (plan):** {U['saat']}  ",
                    f"**İçerik çerçevesi:** {' · '.join(U['konu'])}  ", f"**Anahtar kavramlar:** {', '.join(dict.fromkeys(', '.join(U['ak']).split(', ')))}  ", ""]
        for kod, (ifade, bil) in U["cikti"].items():
            bloklar.append(f"**{kod}.** {ifade}"); bloklar += [f"- {b}" for b in bil]; bloklar.append("")
    (OUT / f"{g}-sinif-din-kulturu-ve-ahlak-bilgisi.md").write_text("\n".join(L + [""] + bloklar).rstrip() + "\n", encoding="utf-8")
    print(f"DKAB {g}: {len(sira)} ünite · {toplam} çıktı")

# ---------------------------------------------------------------- 2) İngilizce 7 (eski program)
xl = glob.glob(str(PK / "ingilizce" / "7" / "*.xlsx"))[0]
ws = openpyxl.load_workbook(xl, data_only=True).worksheets[0]
rows = list(ws.iter_rows(values_only=True))
hi = next(i for i, r in enumerate(rows) if any("LEARNING OUTCOME" in temiz(c).upper() for c in r))
H = [temiz(c).upper() for c in rows[hi]]
c_unit = next(i for i, h in enumerate(H) if "UNIT" in h); c_fn = next(i for i, h in enumerate(H) if "FUNCTIONS" in h)
c_lo = next(i for i, h in enumerate(H) if "LEARNING OUTCOME" in h); c_hr = next(i for i, h in enumerate(H) if "HOURS" in h)
units = {}; sira = []
for r in rows[hi + 1:]:
    # satırlar birleştirilmiş hücreler yüzünden kayabilir: birim adını "N Ad" deseninden bul
    hucre = [temiz(x) for x in r]
    u = next((x for x in hucre if re.match(r"^\d{1,2} [A-Z][A-Za-z ]+$", x)), "")
    if not u: continue
    if u not in units: units[u] = dict(saat=0, fn="", lo=""); sira.append(u)
    U = units[u]
    try: U["saat"] += int(float(r[c_hr]))
    except Exception: pass
    ham = [str(x or "") for x in r]
    lo = next((x for x in ham if re.search(r"E7\.\d+\.[A-Z]+\d", x)), "")
    if lo and not U["lo"]: U["lo"] = lo
    fn = next((x for x in ham if x is not lo and ("—" in x or "Making" in x or "Talking" in x or "Describing" in x)), "")
    if fn and not U["fn"]: U["fn"] = fn
L = ["# 7. Sınıf İngilizce — Kazanım Verisi (2026-2027, ESKİ PROGRAM)", "",
     f"**Kaynak (doğrulanmış):** MEB 2026-2027 İngilizce taslak çerçeve yıllık planı, 7. sınıf sayfası (`{Path(xl).name}`), {bugun} tarihinde aktarıldı "
     "(`tools/kazanim_ozel_uret.py`). Kazanım kodları ve ifadeleri birebir.", "",
     "> **PROGRAM NOTU:** 7. sınıf İngilizce 2026-2027'de okullarda **eski programla** (2018 İngilizce öğretim programı, 10 ünite) okutulur. "
     "ÖDSGM ortak sınav senaryo tablosu ve MEB yıllık planı `E7.<ünite>.<beceri><no>` kodlarını kullanır. TYMM sitesindeki 7. sınıf İngilizce "
     "(ENG.7.x, 8 tema) bu yıl okulda **uygulanmaz**. Terminoloji: \"kazanım\"; süreç bileşeni katmanı yoktur.", "",
     "Beceri kısaltmaları: L dinleme · SI sözlü etkileşim · SP sözlü üretim · R okuma · W yazma.", "",
     "## Ünite tablosu", "", "| Kod/No | Ad | Kazanım | Saat |", "|---|---|---|---|"]
B = []
for u in sira:
    U = units[u]; no, ad = u.split(" ", 1)
    kaz = re.findall(r"(E7\.\d+\.[A-Z]+\d+)\.\s*(.+?)(?=\s*(?:Listening|Spoken Interaction|Spoken Production|Reading|Writing)?\s*E7\.\d+\.[A-Z]+\d+\.|$)", temiz(U["lo"]))
    L.append(f"| ÜNİTE {no} | {ad.upper()} | {len(kaz)} | {U['saat']} |")
    B += [f"## ÜNİTE {no} — {ad.upper()}", "", f"**Ders saati (plan):** {U['saat']}  ", f"**İşlevler ve dil kullanımı:** {temiz(U['fn'])[:900]}  ", ""]
    for k, t in kaz:
        t = re.sub(r"\s*(Listening|Spoken Interaction|Spoken Production|Reading|Writing)\s*$", "", t).strip()
        B.append(f"**{k}.** {t}"); B.append("")
(OUT / "7-sinif-ingilizce.md").write_text("\n".join(L + [""] + B).rstrip() + "\n", encoding="utf-8")
print(f"İngilizce 7 (eski program): {len(sira)} ünite")

# ---------------------------------------------------------------- 3) Türkçe 5 — T.O.5.14
p = OUT / "5-sinif-turkce.md"; t = p.read_text(encoding="utf-8")
if "**T.O.5.14.**" not in t:
    blok = t[t.index("## TEMA 6"):]
    yeni = re.sub(r"(### Okuma\n)", r"\1- **T.O.5.14.** Bilgilendirici metinde metin yapılarından hareketle önemli bilgileri belirlemeye yönelik çözümleme yapabilme "
                  r"*(tam program PDF'i 02.09.2026, 6. Tema; TYMM tema sayfasının çıktı listesinde yok)*\n", blok, count=1)
    p.write_text(t[:t.index("## TEMA 6")] + yeni, encoding="utf-8"); print("Türkçe 5: T.O.5.14 eklendi (Tema 6)")
